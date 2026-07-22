import tempfile
import unittest
from pathlib import Path

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

from core.xbench import filter_xlsx


class FilterXlsxTests(unittest.TestCase):
    def test_preserves_rich_text_highlights_in_kept_rows(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            src = Path(tmp_dir) / "source.xlsx"
            dst = Path(tmp_dir) / "filtered.xlsx"

            wb = openpyxl.Workbook()
            ws = wb.active
            ws["A11"] = "Key Term Mismatch (中文 / Chinese)"
            ws["A12"] = "sample.txt (1)"
            ws["C12"] = "删除此行"
            ws["A13"] = "sample.txt (2)"
            ws["C13"] = CellRichText(
                "保留",
                TextBlock(InlineFont(color="FFFF0000"), "红色标记"),
                "文本",
            )
            ws.row_dimensions[13].height = 88
            wb.save(src)
            wb.close()

            filter_xlsx(str(src), {13}, str(dst))

            filtered = openpyxl.load_workbook(dst, rich_text=True)
            value = filtered.active["C12"].value
            self.assertIsInstance(value, CellRichText)
            self.assertEqual(str(value), "保留红色标记文本")
            red_runs = [
                part
                for part in value
                if isinstance(part, TextBlock)
                and part.font.color is not None
                and part.font.color.rgb == "FFFF0000"
            ]
            self.assertEqual([part.text for part in red_runs], ["红色标记"])
            self.assertEqual(filtered.active.row_dimensions[12].height, 88)
            self.assertIsNone(filtered.active["A13"].value)
            filtered.close()


if __name__ == "__main__":
    unittest.main()
